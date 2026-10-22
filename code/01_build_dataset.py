"""Сборка очищенных и гармонизированных наборов данных для гравитационного анализа.

Вход
----
  * ``Данные.xlsx`` — панель ВКР: 496 пар стран × 2000–2024 гг. (суммарная торговля);
  * открытые зеркала данных Всемирного банка (ВВП в текущих долл. США, население),
    клонированные из github.com/datasets/{gdp,population};
  * CEPII GeoDist (Mayer & Zignago, 2011) — расстояния, общая граница, язык,
    колониальные связи (R-пакет cepiigeodist, github.com/pachadotdev/cepiigeodist).

Выход (каталог ``data/``)
-------------------------
  * ``vkr_panel_renamed.csv|xlsx`` — исходные данные ВКР с переименованными по
    общепринятой нотации переменными (ориентация пар и единицы не изменены);
  * ``gravity_panel.csv|xlsx`` — гармонизированная панель для оценивания;
  * ``countries.csv``, ``membership_timeline.csv`` — справочники;
  * ``data_quality_report.csv`` — перечень найденных дефектов исходных данных.

Запуск:  python code/01_build_dataset.py [--wb-dir DIR] [--geodist FILE]
"""
from __future__ import annotations

import argparse
import itertools
from pathlib import Path

import numpy as np
import pandas as pd

from gravity_common import (COALITION_2014, COALITION_2022, COUNTRIES, DATA, NAME2ISO,
                            ROOT, UKRAINE, YEARS, BRICS_SPAN, EAEU_SPAN, EU_SPAN,
                            NATO_SPAN, WTO_SPAN, both_members, member, partner_group,
                            RU_NAME, SUBGROUP_OF)

VKR_FILE = ROOT / "Данные.xlsx"


# ----------------------------------------------------------------------------
# 1. Исходные данные ВКР
# ----------------------------------------------------------------------------
VKR_RENAME = {
    # позиционно: имена в исходном файле содержат кириллическую «С» и пробелы
    0: "country_i", 1: "country_j", 2: "pair_id", 3: "year", 4: "trade",
    5: "gdppc_i", 6: "gdppc_j", 7: "dist_vkr",
    8: "post2014", 9: "contig", 10: "eu_ij", 11: "brics_ij", 12: "nato_ij",
    13: "post2022",
}


def read_vkr() -> pd.DataFrame:
    raw = pd.read_excel(VKR_FILE)
    assert raw.shape == (12400, 14), raw.shape
    raw.columns = [VKR_RENAME[k] for k in range(raw.shape[1])]
    raw["iso3_i"] = raw["country_i"].map(NAME2ISO)
    raw["iso3_j"] = raw["country_j"].map(NAME2ISO)
    assert raw["iso3_i"].notna().all() and raw["iso3_j"].notna().all()
    return raw


# ----------------------------------------------------------------------------
# 2. Макроданные: ВВП и население
# ----------------------------------------------------------------------------
def load_macro(wb_dir: Path, vkr: pd.DataFrame) -> pd.DataFrame:
    """ВВП (млрд долл.), население (млн), ВВП на душу населения (долл.).

    2000–2023: ВВП — WDI (NY.GDP.MKTP.CD), население — WDI (SP.POP.TOTL);
    2024: ВВП на душу населения из файла ВКР (оценка МВФ WEO, апрель 2025)
    × население WDI.
    """
    gdp = pd.read_csv(wb_dir / "gdp" / "data" / "gdp.csv")
    pop = pd.read_csv(wb_dir / "population" / "data" / "population.csv")
    gdp = gdp.rename(columns={"Country Code": "iso3", "Value": "gdp_usd"})[["iso3", "Year", "gdp_usd"]]
    pop = pop.rename(columns={"Country Code": "iso3", "Value": "pop"})[["iso3", "Year", "pop"]]
    iso_list = list(COUNTRIES)
    idx = pd.MultiIndex.from_product([iso_list, YEARS], names=["iso3", "year"])
    macro = pd.DataFrame(index=idx).reset_index()
    macro = macro.merge(gdp.rename(columns={"Year": "year"}), on=["iso3", "year"], how="left")
    macro = macro.merge(pop.rename(columns={"Year": "year"}), on=["iso3", "year"], how="left")

    # ВВП на душу населения из файла ВКР (для 2024 г.)
    v = pd.concat([
        vkr[["iso3_i", "year", "gdppc_i"]].rename(columns={"iso3_i": "iso3", "gdppc_i": "gdppc_vkr"}),
        vkr[["iso3_j", "year", "gdppc_j"]].rename(columns={"iso3_j": "iso3", "gdppc_j": "gdppc_vkr"}),
    ]).drop_duplicates(["iso3", "year"])
    macro = macro.merge(v, on=["iso3", "year"], how="left")

    macro["gdp_src"] = np.where(macro["gdp_usd"].notna(), "WDI", "")
    miss = macro["gdp_usd"].isna() & macro["gdppc_vkr"].notna() & macro["pop"].notna()
    macro.loc[miss, "gdp_usd"] = macro.loc[miss, "gdppc_vkr"] * macro.loc[miss, "pop"]
    macro.loc[miss, "gdp_src"] = "IMF-WEO x WDI-pop"
    macro["gdp"] = macro["gdp_usd"] / 1e9            # млрд долл.
    macro["pop_m"] = macro["pop"] / 1e6               # млн чел.
    macro["gdppc"] = macro["gdp_usd"] / macro["pop"]  # долл. на душу
    return macro[["iso3", "year", "gdp", "pop_m", "gdppc", "gdp_src"]]


# ----------------------------------------------------------------------------
# 3. CEPII GeoDist
# ----------------------------------------------------------------------------
def load_geodist(path: Path) -> pd.DataFrame:
    import pyreadr  # pip install pyreadr
    d = pyreadr.read_r(str(path))["dist_cepii"]
    for c in ["distw", "distwces"]:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    keep = ["iso_o", "iso_d", "contig", "comlang_off", "comlang_ethno", "colony", "comcol",
            "curcol", "col45", "smctry", "dist", "distcap", "distw", "distwces"]
    return d[keep]


# ----------------------------------------------------------------------------
# 4. Сборка гармонизированной панели
# ----------------------------------------------------------------------------
def build_panel(vkr: pd.DataFrame, macro: pd.DataFrame, geo: pd.DataFrame) -> pd.DataFrame:
    # каноническая ориентация пары: iso3_i < iso3_j (торговля симметрична)
    trade = vkr[["iso3_i", "iso3_j", "year", "trade", "pair_id"]].copy()
    swap = trade["iso3_i"] > trade["iso3_j"]
    trade.loc[swap, ["iso3_i", "iso3_j"]] = trade.loc[swap, ["iso3_j", "iso3_i"]].values
    trade["trade"] = trade["trade"] / 1e3  # тыс. долл. -> млн долл.

    pairs = trade[["iso3_i", "iso3_j"]].drop_duplicates().reset_index(drop=True)
    pairs["pair_id_h"] = np.arange(1, len(pairs) + 1)
    trade = trade.merge(pairs, on=["iso3_i", "iso3_j"], how="left")

    p = trade.copy()
    for side in ("i", "j"):
        m = macro.rename(columns={"iso3": f"iso3_{side}", "gdp": f"gdp_{side}", "pop_m": f"pop_{side}",
                                  "gdppc": f"gdppc_{side}", "gdp_src": f"gdp_src_{side}"})
        p = p.merge(m, on=[f"iso3_{side}", "year"], how="left")
    p["gdp_src"] = p["gdp_src_i"]
    p = p.drop(columns=["gdp_src_i", "gdp_src_j"])

    # CEPII: ищем пару в обеих ориентациях
    g1 = geo.rename(columns={"iso_o": "iso3_i", "iso_d": "iso3_j"})
    g2 = geo.rename(columns={"iso_o": "iso3_j", "iso_d": "iso3_i"})
    gg = pd.concat([g1, g2]).drop_duplicates(["iso3_i", "iso3_j"])
    p = p.merge(gg.drop(columns=["dist", "distwces"]), on=["iso3_i", "iso3_j"], how="left")
    p = p.rename(columns={"distw": "dist_w", "distcap": "dist_cap"})
    p["dist"] = p["dist_w"]  # основная мера расстояния: население-взвешенное (CEPII distw), км

    # членство в объединениях (парные индикаторы «обе страны — члены»)
    for name, span in [("eu_ij", EU_SPAN), ("nato_ij", NATO_SPAN), ("brics_ij", BRICS_SPAN),
                       ("eaeu_ij", EAEU_SPAN), ("wto_ij", WTO_SPAN)]:
        p[name] = [both_members(span, a, b, y) for a, b, y in zip(p.iso3_i, p.iso3_j, p.year)]

    # российские пары и санкционные индикаторы
    p["rus_pair"] = ((p.iso3_i == "RUS") | (p.iso3_j == "RUS")).astype(int)
    p["partner"] = np.where(p.iso3_i == "RUS", p.iso3_j, np.where(p.iso3_j == "RUS", p.iso3_i, ""))
    p["partner_group"] = [partner_group(x) if x else "" for x in p.partner]
    p["partner_subgroup"] = [SUBGROUP_OF.get(x, "") for x in p.partner]
    p["coal_rus_ij"] = (p.rus_pair.eq(1) & p.partner.isin(COALITION_2022)).astype(int)  # группа «обработки»
    p["sanc14_ij"] = (p.rus_pair.eq(1) & p.partner.isin(COALITION_2014) & (p.year >= 2014)).astype(int)
    p["sanc22_ij"] = (p.rus_pair.eq(1) & p.partner.isin(COALITION_2022) & (p.year >= 2022)).astype(int)
    p["ukr14_ij"] = (p.rus_pair.eq(1) & p.partner.eq(UKRAINE) & (p.year >= 2014)).astype(int)
    p["ukr22_ij"] = (p.rus_pair.eq(1) & p.partner.eq(UKRAINE) & (p.year >= 2022)).astype(int)

    # индикаторы качества данных
    p = p.sort_values(["pair_id_h", "year"]).reset_index(drop=True)
    prev = p.groupby("pair_id_h")["trade"].shift(1)
    p["flag_carry"] = (p["trade"] == prev).astype(int)  # значение «протянуто» с прошлого года
    p["trade_src"] = "VKR (UN Comtrade, суммарная торговля)"

    cols = ["pair_id_h", "year", "iso3_i", "iso3_j", "trade",
            "gdp_i", "gdp_j", "gdppc_i", "gdppc_j", "pop_i", "pop_j", "gdp_src",
            "dist", "dist_cap", "contig", "comlang_off", "comlang_ethno", "colony", "comcol", "col45", "smctry",
            "eu_ij", "nato_ij", "brics_ij", "eaeu_ij", "wto_ij",
            "rus_pair", "partner", "partner_group", "partner_subgroup",
            "coal_rus_ij", "sanc14_ij", "sanc22_ij", "ukr14_ij", "ukr22_ij",
            "flag_carry", "trade_src"]
    p = p[cols].rename(columns={"pair_id_h": "pair_id"})
    # названия стран для удобства чтения
    p.insert(4, "country_j", p["iso3_j"].map(lambda x: COUNTRIES[x][0]))
    p.insert(4, "country_i", p["iso3_i"].map(lambda x: COUNTRIES[x][0]))
    return p


# ----------------------------------------------------------------------------
# 5. Отчёт о качестве исходных данных
# ----------------------------------------------------------------------------
def quality_report(vkr: pd.DataFrame, geo: pd.DataFrame) -> pd.DataFrame:
    rows = []
    # (a) «санкционные» дамми ВКР — временные индикаторы
    for var, y0 in [("post2014", 2014), ("post2022", 2022)]:
        share = vkr.groupby("year")[var].mean()
        ok = bool((share.loc[share.index >= y0] == 1).all() and (share.loc[share.index < y0] == 0).all())
        rows.append(("санкционный дамми — чисто временной", var,
                     f"равен 1 для ВСЕХ 496 пар при year>={y0} (временной дамми, не парный эффект): {ok}",
                     int(vkr[var].sum())))
    # (b) БРИКС 2024: ошибки кодировки
    s = vkr[(vkr.year == 2024) & (vkr.brics_ij == 1)]
    flagged = set(map(frozenset, zip(s.iso3_i, s.iso3_j)))
    members = {"BRA", "CHN", "IND", "RUS", "ZAF", "EGY", "ARE"}
    correct = {frozenset(c) for c in itertools.combinations(sorted(members), 2)}
    for pair in sorted(flagged - correct, key=lambda x: sorted(x)):
        rows.append(("ошибка кодировки БРИКС (лишняя пара)", "brics_ij", "2024: " + "–".join(sorted(pair)), 1))
    for pair in sorted(correct - flagged, key=lambda x: sorted(x)):
        rows.append(("ошибка кодировки БРИКС (пропущенная пара)", "brics_ij", "2024: " + "–".join(sorted(pair)), 1))
    # (c) НАТО 2000–2003
    s = vkr[(vkr.year == 2000) & (vkr.nato_ij == 1)]
    flagged = set(map(frozenset, zip(s.iso3_i, s.iso3_j)))
    mem = {"USA", "CAN", "GBR", "FRA", "DEU", "ITA", "NOR", "POL", "TUR"}
    correct = {frozenset(c) for c in itertools.combinations(sorted(mem), 2)}
    for pair in sorted(correct - flagged, key=lambda x: sorted(x)):
        rows.append(("ошибка кодировки НАТО (пропущенная пара)", "nato_ij", "2000–2003: " + "–".join(sorted(pair)), 4))
    # (d) расстояния: расхождение с CEPII distcap > 10%
    pairs = vkr[["iso3_i", "iso3_j", "dist_vkr"]].drop_duplicates()
    gg = pd.concat([geo.rename(columns={"iso_o": "iso3_i", "iso_d": "iso3_j"}),
                    geo.rename(columns={"iso_o": "iso3_j", "iso_d": "iso3_i"})]).drop_duplicates(["iso3_i", "iso3_j"])
    m = pairs.merge(gg, on=["iso3_i", "iso3_j"])
    m["ratio"] = m["dist_vkr"] / m["distcap"]
    bad = m[(m.ratio < 0.9) | (m.ratio > 1.1)]
    rows.append(("расстояния ВКР расходятся с CEPII distcap более чем на 10%", "dist_vkr",
                 f"{len(bad)} из {len(m)} пар; максимум: " +
                 "; ".join(f"{a}-{b} {r:.2f}" for a, b, r in
                           bad.reindex(bad.ratio.sub(1).abs().sort_values(ascending=False).index)[["iso3_i", "iso3_j", "ratio"]].head(5).values),
                 len(bad)))
    # (e) протягивание значений торговли
    w = vkr.pivot(index=["iso3_i", "iso3_j"], columns="year", values="trade")
    for y in range(2001, 2025):
        n = int((w[y] == w[y - 1]).sum())
        if n:
            rows.append(("значение торговли в году t совпадает с t-1 (протянуто)", "trade", f"год {y}", n))
    # (f) обвалы в 2022 г. в парах без России
    lw = np.log(w)
    d22 = (lw[2022] - lw[2021])
    idx = d22.index.to_frame(index=False)
    nr = ~((idx.iso3_i == "RUS") | (idx.iso3_j == "RUS")).values
    share = float((d22[nr].abs() > 1.0).mean())
    ref = float(((lw[2019] - lw[2018])[nr].abs() > 1.0).mean())
    rows.append(("аномально частые скачки |Δln trade|>1 в парах без России", "trade",
                 f"2022 г.: {share:.1%} пар против {ref:.1%} в 2019 г.", int((d22[nr].abs() > 1.0).sum())))
    # (g) ВВП на душу населения: смена источника после 2021 г.
    rows.append(("смена источника ВВП на душу населения", "gdppc_i/gdppc_j",
                 "2000–2021: Всемирный банк (WDI); 2022–2024: МВФ WEO (3 знака после запятой)", 0))
    return pd.DataFrame(rows, columns=["issue", "variable", "details", "n_obs_or_pairs"])


# ----------------------------------------------------------------------------
# 6. Экспорт в Excel и словарь данных
# ----------------------------------------------------------------------------
def _format_sheet(ws, widths=None, freeze="A2") -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, start=1):
        w = (widths or {}).get(i)
        if w is None:
            m = max((len(str(c.value)) if c.value is not None else 0) for c in list(col)[:200])
            w = min(max(10, m + 2), 60)
        ws.column_dimensions[get_column_letter(i)].width = w


def export_excel(vkr_cols: pd.DataFrame, panel: pd.DataFrame, cdf: pd.DataFrame, mt: pd.DataFrame,
                 qr: pd.DataFrame) -> None:
    from data_dictionary import PANEL_VARS
    dd = pd.DataFrame(PANEL_VARS, columns=["variable", "notation", "name_in_VKR", "units", "description"])
    with pd.ExcelWriter(DATA / "gravity_panel.xlsx", engine="openpyxl") as xw:
        panel.to_excel(xw, sheet_name="panel", index=False)
        dd.to_excel(xw, sheet_name="dictionary", index=False)
        cdf.to_excel(xw, sheet_name="countries", index=False)
        mt.to_excel(xw, sheet_name="membership", index=False)
        qr.to_excel(xw, sheet_name="data_quality", index=False)
        for ws in xw.book.worksheets:
            _format_sheet(ws, widths={5: 24} if ws.title == "dictionary" else None)
        xw.book["dictionary"].column_dimensions["E"].width = 110
    with pd.ExcelWriter(DATA / "vkr_panel_renamed.xlsx", engine="openpyxl") as xw:
        vkr_cols.to_excel(xw, sheet_name="panel", index=False)
        dd.to_excel(xw, sheet_name="dictionary", index=False)
        for ws in xw.book.worksheets:
            _format_sheet(ws)
        xw.book["dictionary"].column_dimensions["E"].width = 110
    # Markdown
    lines = ["# Словарь данных", "",
             "Файлы: `gravity_panel.csv|xlsx` (гармонизированная панель), `vkr_panel_renamed.csv|xlsx` "
             "(исходные данные ВКР с переименованными переменными).", "",
             "| Переменная | Обозначение | Имя в ВКР | Единицы / значения | Описание |",
             "|---|---|---|---|---|"]
    for r in PANEL_VARS:
        lines.append("| " + " | ".join(f"`{x}`" if k == 0 else x for k, x in enumerate(r)) + " |")
    (DATA / "DATA_DICTIONARY.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--wb-dir", default="/home/user/datasets", help="каталог с клонами datasets/gdp и datasets/population")
    ap.add_argument("--geodist", default="/home/user/pachadotdev/cepiigeodist/data/dist_cepii.rda")
    args = ap.parse_args()

    DATA.mkdir(exist_ok=True)
    vkr = read_vkr()
    macro = load_macro(Path(args.wb_dir), vkr)
    geo = load_geodist(Path(args.geodist))

    # --- набор 1: исходные данные ВКР, переименованные ---
    cols1 = ["pair_id", "year", "iso3_i", "iso3_j", "country_i", "country_j", "trade", "gdppc_i", "gdppc_j",
             "dist_vkr", "contig", "post2014", "post2022", "eu_ij", "brics_ij", "nato_ij"]
    vkr[cols1].to_csv(DATA / "vkr_panel_renamed.csv", index=False)

    # --- набор 2: гармонизированная панель ---
    panel = build_panel(vkr, macro, geo)
    panel.to_csv(DATA / "gravity_panel.csv", index=False)

    # --- справочники ---
    cdf = pd.DataFrame([(k, v[0], v[1], v[2], int(k in COALITION_2014), int(k in COALITION_2022))
                        for k, v in COUNTRIES.items()],
                       columns=["iso3", "country", "country_ru", "region", "sanctions_2014", "sanctions_2022"])
    cdf.to_csv(DATA / "countries.csv", index=False)
    mt = pd.DataFrame([(iso, y, member(EU_SPAN, iso, y), member(NATO_SPAN, iso, y), member(BRICS_SPAN, iso, y),
                        member(EAEU_SPAN, iso, y), member(WTO_SPAN, iso, y))
                       for iso in COUNTRIES for y in YEARS],
                      columns=["iso3", "year", "eu", "nato", "brics", "eaeu", "wto"])
    mt.to_csv(DATA / "membership_timeline.csv", index=False)
    macro.to_csv(DATA / "macro_country_year.csv", index=False)

    qr = quality_report(vkr, geo)
    qr.to_csv(DATA / "data_quality_report.csv", index=False)

    export_excel(vkr[cols1], panel, cdf, mt, qr)

    print("vkr_panel_renamed:", vkr[cols1].shape, "| gravity_panel:", panel.shape)
    print(panel.isna().sum()[panel.isna().sum() > 0])
    print(qr.to_string())


if __name__ == "__main__":
    main()
